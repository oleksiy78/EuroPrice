#!/usr/bin/env python3
"""Update EuroPrice fuel prices from the European Commission's official XLSX."""
import json
import re
from datetime import datetime, timezone
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urljoin

import requests
from openpyxl import load_workbook

BULLETIN_URL = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
PRICE_FILE = Path("prices.json")
HEADERS = {
    "User-Agent": "EuroPrice/1.0 (fuel price updater; https://github.com/oleksiy78/EuroPrice)",
    "Accept": "text/html,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
}

COUNTRIES = {
    "germany": "DE",
    "spain": "ES",
    "france": "FR",
    "italy": "IT",
    "austria": "AT",
    "croatia": "HR",
    "portugal": "PT",
    "greece": "GR",
    "czechia": "CZ",
    "czech republic": "CZ",
    "poland": "PL",
}
PRICE_LABELS = {
    "petrol": ("euro-super 95", "eurosuper 95", "euro super 95"),
    "diesel": ("automotive gas oil", "diesel", "gas oil automobile"),
}


class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self._href = None
        self._text = []

    def handle_starttag(self, tag, attrs):
        if tag.lower() == "a":
            self._href = dict(attrs).get("href")
            self._text = []

    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag):
        if tag.lower() == "a" and self._href is not None:
            self.links.append((self._href, " ".join(self._text).strip()))
            self._href = None
            self._text = []


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip().lower()


def find_xlsx_url(page_url, html):
    parser = LinkParser()
    parser.feed(html)
    candidates = []
    for href, label in parser.links:
        full_url = urljoin(page_url, href or "")
        decoded = unquote(full_url).lower()
        if "/document/download/" in decoded and ".xlsx" in decoded and "with" in decoded and "taxes" in decoded:
            candidates.append(full_url)
    if not candidates:
        raise RuntimeError("Official EC page did not expose the 'Prices with taxes' XLSX download link.")
    return candidates[0]


def parse_number(value):
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value or "").replace("\xa0", "").replace(" ", "").strip()
        if not text:
            raise ValueError("Empty price cell")
        # EC workbooks usually use decimal commas or dot decimals.
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
        match = re.search(r"-?\d+(?:\.\d+)?", text)
        if not match:
            raise ValueError(f"Cannot parse price: {value!r}")
        number = float(match.group(0))
    # The bulletin's published unit is EUR per 1,000 litres; accept €/L too.
    if number > 10:
        number /= 1000.0
    if not 0.5 <= number <= 4.0:
        raise ValueError(f"Fuel price outside safety range after unit conversion: {value!r} -> {number}")
    return round(number, 4)


def parse_workbook(content):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    found = {}
    bulletin_date = None
    country_aliases = set(COUNTRIES)
    date_patterns = (
        r"(\d{1,2})[./ -](\d{1,2})[./ -](\d{4})",
        r"(\d{4})-(\d{2})-(\d{2})",
    )

    for sheet in workbook.worksheets:
        grid = list(sheet.iter_rows(values_only=True))
        # Recover the bulletin date from a title/caption cell if present.
        for row in grid[:12]:
            for cell in row:
                if not isinstance(cell, str):
                    continue
                text = norm(cell)
                if "prices in force on" in text or "preisstand" in text or "prix en vigueur au" in text:
                    for pattern in date_patterns:
                        match = re.search(pattern, cell)
                        if match:
                            parts = match.groups()
                            try:
                                if len(parts[0]) == 4:
                                    dt = datetime(int(parts[0]), int(parts[1]), int(parts[2]))
                                else:
                                    dt = datetime(int(parts[2]), int(parts[1]), int(parts[0]))
                                bulletin_date = dt.strftime("%d %b %Y")
                                break
                            except ValueError:
                                pass

        # Header labels can span more than one row in spreadsheet exports.
        petrol_col = diesel_col = None
        header_end = None
        max_cols = max((len(row) for row in grid), default=0)
        for r in range(min(25, len(grid))):
            for c in range(max_cols):
                combined = " ".join(
                    norm(grid[rr][c]) for rr in range(r, min(r + 3, len(grid)))
                    if c < len(grid[rr]) and grid[rr][c] is not None
                )
                if petrol_col is None and any(label in combined for label in PRICE_LABELS["petrol"]):
                    petrol_col = c
                if diesel_col is None and any(label in combined for label in PRICE_LABELS["diesel"]):
                    diesel_col = c
                if petrol_col is not None and diesel_col is not None:
                    header_end = r
                    break
            if petrol_col is not None and diesel_col is not None:
                break
        if petrol_col is None or diesel_col is None:
            continue

        for row in grid[header_end or 0:]:
            if not row:
                continue
            country_code = None
            for cell in row:
                if isinstance(cell, str) and norm(cell) in country_aliases:
                    country_code = COUNTRIES[norm(cell)]
                    break
            if not country_code or country_code in found:
                continue
            try:
                petrol = parse_number(row[petrol_col])
                diesel = parse_number(row[diesel_col])
            except (IndexError, ValueError, TypeError):
                continue
            found[country_code] = (petrol, diesel)
        if set(found) == set(COUNTRIES.values()):
            break

    workbook.close()
    missing = set(COUNTRIES.values()) - set(found)
    if missing:
        raise RuntimeError(
            "Could not safely parse official EC XLSX; missing country codes: "
            + ", ".join(sorted(missing))
        )
    return found, bulletin_date


def main():
    session = requests.Session()
    page_response = session.get(BULLETIN_URL, headers=HEADERS, timeout=30)
    page_response.raise_for_status()
    if "text/html" not in page_response.headers.get("Content-Type", "").lower():
        raise RuntimeError("Official EC bulletin page did not return HTML.")
    xlsx_url = find_xlsx_url(BULLETIN_URL, page_response.text)

    file_response = session.get(xlsx_url, headers=HEADERS, timeout=30)
    file_response.raise_for_status()
    content_type = file_response.headers.get("Content-Type", "").lower()
    if "spreadsheet" not in content_type and not file_response.content.startswith(b"PK"):
        raise RuntimeError(
            f"EC XLSX download did not look like an Excel file (Content-Type: {content_type})."
        )

    rows, bulletin_date = parse_workbook(file_response.content)
    if not bulletin_date:
        # The official link filename normally includes the bulletin date.
        filename = unquote(xlsx_url)
        match = re.search(r"(\d{4})-(\d{2})-(\d{2})\.xlsx", filename)
        if match:
            bulletin_date = datetime(
                int(match.group(1)), int(match.group(2)), int(match.group(3))
            ).strftime("%d %b %Y")
    if not bulletin_date:
        raise RuntimeError("Parsed prices, but could not verify the bulletin date; refusing to update.")

    price_file = json.loads(PRICE_FILE.read_text(encoding="utf-8"))
    for code, (petrol, diesel) in rows.items():
        price_file["prices"][code]["fuel"] = {
            "price": petrol,
            "unit": "€/L",
            "source": "European Commission — Weekly Oil Bulletin",
            "date": bulletin_date,
            "quality": "national average incl. taxes",
        }
        price_file["prices"][code]["diesel"] = {
            "price": diesel,
            "unit": "€/L",
            "source": "European Commission — Weekly Oil Bulletin",
            "date": bulletin_date,
            "quality": "national average incl. taxes",
        }

    today = datetime.now(timezone.utc).strftime("%d.%m.%Y")
    price_file["meta"]["updated"] = today
    price_file["meta"]["fuelUpdated"] = today
    PRICE_FILE.write_text(
        json.dumps(price_file, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(f"Updated fuel prices from European Commission bulletin dated {bulletin_date}.")
    for code in sorted(rows):
        print(f"{code}: Super 95 €{rows[code][0]:.3f}/L, Diesel €{rows[code][1]:.3f}/L")


if __name__ == "__main__":
    main()
